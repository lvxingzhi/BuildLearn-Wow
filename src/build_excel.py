#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_excel.py — 汇总成 Excel 报告, 归档到 runs/<日期>-<职业>-<专精>/

输入: data/clean/spells.json + data/clean/talents.json + cache/wowhead_cache.json
输出: runs/<slug>/
        ├── 分析报告.xlsx   Tab1 施法次数 / Tab2 天赋 (图标直接嵌入单元格)
        ├── spells.json     本次 run 的清洗产物副本 (存档)
        ├── talents.json
        └── meta.json       职业/专精/天赋码/条数等摘要

用法:
    python3 src/build_excel.py                    # 自动按 日期-职业-专精 建目录
    python3 src/build_excel.py --name my-run      # 自定义目录名
    python3 src/build_excel.py --out /tmp/a.xlsx  # 直接指定输出文件 (不归档)
"""

import argparse
import json
import shutil
import sys
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config  # noqa: E402

ICON_BASE = "https://assets.rpglogs.cn/img/warcraft/abilities/"  # WCL 图标 CDN

# ── 样式 ────────────────────────────────────────────────────────────────
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="4472C4")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")

GROUP_CN = {"class": "职业", "spec": "专精", "hero": "英雄"}
GROUP_FILL = {"class": "D9E1F2", "spec": "E2EFDA", "hero": "FCE4D6"}


def style_header(ws, ncols, row=1):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fetch_icon(url, cache_dir=None):
    """下载图标到本地缓存, 返回本地路径; 失败返回 None."""
    if not url:
        return None
    d = Path(cache_dir) if cache_dir else config.ICON_DIR
    d.mkdir(parents=True, exist_ok=True)
    name = url.rsplit("/", 1)[-1]
    if not name:
        return None
    local = d / name
    if local.exists():
        return str(local)
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=8) as r:
            local.write_bytes(r.read())
        return str(local)
    except Exception:
        return None


def embed_icon(ws, cell, spell_id, info):
    """把一个 24x24 图标嵌到指定单元格."""
    icon = (info.get("icon") or "").strip()
    if not icon:
        return False
    url = icon if icon.startswith("http") else "{}{}.jpg".format(ICON_BASE, icon)
    local = fetch_icon(url)
    if not local:
        return False
    try:
        img = XLImage(local)
        img.width = 24
        img.height = 24
        ws.add_image(img, cell)
        return True
    except Exception:
        return False


def row_height(desc, base=14, per=45, width=40, lo=28, hi=140):
    """按描述长度估算行高."""
    return max(lo, min(hi, base + len(desc or "") // width * per))


def build_sheet_casts(wb, spells, wh):
    ws = wb.active
    ws.title = "施法次数"
    headers = ["排名", "图标", "技能名", "Spell ID", "施法次数", "占比%", "CPM",
               "Uptime%", "法力/消耗", "射程", "施法时间", "技能描述"]
    ws.append(headers)
    style_header(ws, len(headers))

    for i, s in enumerate(spells, 1):
        info = wh(s["id"])
        ws.append([i, "", info.get("name") or s["name"], s["id"],
                   s["count"], s["pct"], s["cpm"],
                   s["uptime_pct"] if s["uptime_pct"] is not None else "-",
                   info.get("cost", ""), info.get("range", ""),
                   info.get("cast_time", ""), info.get("desc", "")])
        r = ws.max_row
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
        ws.cell(row=r, column=len(headers)).alignment = WRAP_TOP
        ws.row_dimensions[r].height = row_height(info.get("desc", ""))
        embed_icon(ws, "B{}".format(r), s["id"], info)

    # 合计行
    total = sum(s["count"] for s in spells)
    ws.append(["", "", "合计", "", total, 100.0, "", "", "", "", "", ""])
    r = ws.max_row
    for col in range(1, len(headers) + 1):
        ws.cell(row=r, column=col).font = Font(bold=True)
        ws.cell(row=r, column=col).border = BORDER

    set_widths(ws, [6, 8, 16, 10, 10, 8, 8, 10, 14, 12, 12, 60])
    ws.freeze_panes = "C2"
    return ws


def build_sheet_talents(wb, talents, wh):
    ws = wb.create_sheet("天赋")
    p = talents["player"]
    hero_name = next((h["name"] for h in talents["heroTrees"]
                      if h["id"] == p["heroSpecId"]), p["heroSpecId"])

    ws.append(["职业", p["className"], "专精", p["specName"], "英雄天赋", hero_name])
    ws.append(["天赋码", talents["talentCode"], "", "", "", ""])
    ws.append([])

    headers = ["分组", "图标", "天赋名", "Spell ID", "类型", "选择/等级",
               "法力/消耗", "射程", "施法时间", "天赋描述"]
    ws.append(headers)
    header_row = ws.max_row
    style_header(ws, len(headers), row=header_row)

    for i, t in enumerate(talents["selectedTalents"]):
        info = wh(t["spellId"])
        ws.append([GROUP_CN.get(t["treeType"], t["treeType"]), "",
                   info.get("name") or t["name"], t["spellId"], t["type"],
                   t["ranksChosen"] if t["ranksChosen"] is not None else "-",
                   info.get("cost", ""), info.get("range", ""),
                   info.get("cast_time", ""), info.get("desc", "")])
        r = ws.max_row
        fill = PatternFill("solid", fgColor=GROUP_FILL.get(t["treeType"], "FFFFFF"))
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).fill = fill
            ws.cell(row=r, column=col).border = BORDER
        ws.cell(row=r, column=len(headers)).alignment = WRAP_TOP
        ws.row_dimensions[r].height = row_height(info.get("desc", ""), per=14, width=45, hi=140)
        embed_icon(ws, "B{}".format(r), t["spellId"], info)

    set_widths(ws, [8, 8, 24, 10, 10, 10, 14, 12, 12, 65])
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3).coordinate
    return ws, hero_name


def main():
    ap = argparse.ArgumentParser(description="生成 Excel 报告")
    ap.add_argument("--spells", help="默认 data/clean/spells.json")
    ap.add_argument("--talents", help="默认 data/clean/talents.json")
    ap.add_argument("--cache", help="默认 cache/wowhead_cache.json")
    ap.add_argument("--name", help="run 目录名 (默认 日期-职业-专精)")
    ap.add_argument("--out", help="直接指定 xlsx 路径, 不走 runs/ 归档")
    args = ap.parse_args()

    config.ensure_dirs()
    spells_path = config.require(
        Path(args.spells) if args.spells else config.CLEAN_DIR / "spells.json",
        "先跑 python3 src/cleanse_spells.py")
    talents_path = config.require(
        Path(args.talents) if args.talents else config.CLEAN_DIR / "talents.json",
        "先跑 python3 src/cleanse_talents.py")
    cache_path = config.require(
        Path(args.cache) if args.cache else config.CACHE_DIR / "wowhead_cache.json",
        "先跑 python3 src/fetch_wowhead.py")

    spells = json.loads(spells_path.read_text("utf-8"))
    talents = json.loads(talents_path.read_text("utf-8"))
    WH = json.loads(cache_path.read_text("utf-8"))

    def wh(sid):
        """取 Wowhead 缓存, sid 可能是 int/str."""
        return WH.get(str(sid), WH.get(sid, {}))

    wb = Workbook()
    build_sheet_casts(wb, spells, wh)
    _, hero_name = build_sheet_talents(wb, talents, wh)

    # ── 输出 ──
    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        run_dir = None
    else:
        run_dir = config.make_run_dir(talents["player"], args.name)
        out = run_dir / config.REPORT_XLSX
        # 归档本次 run 的清洗产物, 让目录自包含
        for src in (spells_path, talents_path):
            shutil.copy2(src, run_dir / src.name)
        (run_dir / "meta.json").write_text(json.dumps({
            "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "player": talents["player"],
            "heroName": hero_name,
            "talentCode": talents["talentCode"],
            "spellCount": len(spells),
            "talentCount": len(talents["selectedTalents"]),
            "spellsFrom": str(config.rel(spells_path)),
            "talentsFrom": str(config.rel(talents_path)),
        }, ensure_ascii=False, indent=2), "utf-8")

    wb.save(out)

    print("✅ 已生成 {}".format(config.rel(out)))
    print("   施法次数 tab: {} 条技能 (含图标/描述)".format(len(spells)))
    print("   天赋 tab:     {} 个天赋 (含图标/描述)".format(len(talents["selectedTalents"])))
    if run_dir:
        print("   归档目录:     {}".format(config.rel(run_dir)))
        print("   图标缓存:     {}".format(config.rel(config.ICON_DIR)))


if __name__ == "__main__":
    main()
